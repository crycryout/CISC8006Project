#!/usr/bin/env python3
"""Check final PDF/PPTX contents against measured statistics and source images."""
import csv
import hashlib
import json
from pathlib import Path
import re
import subprocess
import sys

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from src.io_utils import file_hash, now, write_json


def require(condition, message):
    if not condition:
        raise ValueError(message)


def measured_tokens(value):
    return [f"{value['paired_macro_delta_nll']:.6f}",
            *(f"{number:.6f}" for number in value['bootstrap_ci95']),
            value['claim_status']]


def main():
    completion = json.loads((ROOT / 'results/completion.json').read_text())
    require(all(completion.get(stage, {}).get('state') == 'complete'
                for stage in ('reproduction', 'pilots', 'improvement')),
            'Complete the registered science before checking final deliverables')
    controls = completion.get('ablations', {})
    require(len(controls) == 2 and all(row.get('state') == 'complete' for row in controls.values()),
            'Both selected-method controls must be complete')
    paths = ['report/report.md', 'report/report.pdf', 'presentation/defense.pptx']
    markdown = (ROOT / paths[0]).read_text()
    pdf_text = subprocess.check_output(['pdftotext', str(ROOT / paths[1]), '-'], text=True)
    pdf_text = pdf_text.replace('\u2212', '-').replace('\u2013', '-')
    pdf_info = subprocess.check_output(['pdfinfo', str(ROOT / paths[1])], text=True)
    pages = int(re.search(r'^Pages:\s+(\d+)', pdf_info, re.MULTILINE).group(1))
    require(pages > 0 and len(pdf_text.strip()) > 1000, 'Final PDF has no readable report content')
    presentation = Presentation(ROOT / paths[2])
    require(len(presentation.slides) == 12, 'Expected 12 slides in the delivered deck')
    slide_texts = []
    slide_checks = []
    image_digests = set()
    for index, slide in enumerate(presentation.slides, 1):
        visible = '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)
        require(slide.has_notes_slide and len(slide.notes_slide.notes_text_frame.text.strip()) > 40,
                f'Slide {index} has no substantive speaker notes')
        for shape in slide.shapes:
            require(shape.left >= 0 and shape.top >= 0 and shape.width > 0 and shape.height > 0
                    and shape.left + shape.width <= presentation.slide_width
                    and shape.top + shape.height <= presentation.slide_height,
                    f'Slide {index} has a shape outside the slide bounds')
            if hasattr(shape, 'image'):
                image_digests.add(hashlib.sha256(shape.image.blob).hexdigest())
        slide_texts.append(visible)
        slide_checks.append(dict(slide=index,notes_characters=len(slide.notes_slide.notes_text_frame.text),
                                 all_shapes_within_bounds=True))
    comparisons = []
    for stage, slide_index in [('reproduction', 4), ('improvement', 8),
                               *(('ablations/' + name, 9) for name in controls)]:
        value = json.loads((ROOT / 'results' / stage / 'summary.json').read_text())
        for token in measured_tokens(value):
            require(token in markdown and token in pdf_text, f'{stage}: report lacks measured {token}')
            require(token in slide_texts[slide_index], f'{stage}: visible slide lacks measured {token}')
        comparisons.append(dict(stage=stage,visible_slide=slide_index+1,metrics_match=True,
                                summary_sha256=file_hash(ROOT / 'results' / stage / 'summary.json')))
    with (ROOT / 'results/pilots/comparison.csv').open() as stream:
        pilot_rows = list(csv.DictReader(stream))
    require(len(pilot_rows) == 2, 'Both pilot candidates must be present')
    for row in pilot_rows:
        for token in [row['method'], f"{float(row['paired_macro_delta_nll']):.6f}"]:
            require(token in markdown and token in pdf_text and token in slide_texts[7],
                    'Pilot content differs from its measured comparison')
    expected_figures = ['figures/reproduction_per_book_delta.png',
                        'figures/diagnostic_nll_vs_position.png',
                        'figures/improvement_quality_cost.png']
    for relative in expected_figures:
        require(file_hash(ROOT / relative) in image_digests, 'Deck image differs from source: ' + relative)
    require('Complete paired results are still pending' not in markdown
            and 'Recorded selection: None' not in markdown,
            'Final report retains an incomplete scientific placeholder')
    receipt = dict(checked_at=now(),status='pass',
                   scope='Agent content/structure verification; no Office visual rendering or human rehearsal claimed',
                   verification_script_sha256=file_hash(Path(__file__)),
                   pdf_pages=pages,deck_slides=12,comparisons=comparisons,pilots_present=True,
                   embedded_figures=[dict(path=relative,sha256=file_hash(ROOT/relative)) for relative in expected_figures],
                   slides=slide_checks,
                   artifacts=[dict(path=relative,sha256=file_hash(ROOT/relative)) for relative in paths])
    write_json(ROOT / 'environment/verification/deliverable-verification.json', receipt)
    print(json.dumps(dict(status='pass',pdf_pages=pages,deck_slides=12,measured_comparisons=len(comparisons))))


if __name__ == '__main__':
    main()
