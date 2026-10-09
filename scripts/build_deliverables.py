#!/usr/bin/env python3
"""Generate an honest report/deck snapshot from existing measured evidence only."""
import argparse
import csv
import html
import json
from pathlib import Path
import sys
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from src.io_utils import file_hash, now, write_json
from src.run_store import spent


def read_optional(relative):
    path=ROOT/relative
    return json.loads(path.read_text()) if path.exists() else None


def paired_text(value):
    if not value: return "Complete paired results are still pending; no verdict is assigned."
    low,high=value["bootstrap_ci95"]
    return (f"Equal-book candidate-minus-baseline NLL difference {value['paired_macro_delta_nll']:.6f} nats/token; "
            f"book-cluster 95% CI [{low:.6f}, {high:.6f}]; verdict {value['claim_status']}. "
            f"{value['n_books']} books, actual model seeds {value['seeds']}. "
            f"Baseline macro NLL {value['baseline']['macro_book_nll']:.6f}, candidate {value['candidate']['macro_book_nll']:.6f}; "
            f"secondary micro PPL {value['baseline']['micro_token_ppl']:.4f} versus {value['candidate']['micro_token_ppl']:.4f}.")


def selected_cases(selected):
    paired=read_optional("results/improvement/paired_result.json")
    if not paired: return "Full-test per-book cases remain pending."
    largest=sorted(paired["per_book"],key=lambda row:row["delta"],reverse=True)[:2]
    text="Largest selected-minus-fixed-four book deltas: "+"; ".join(f"{row['book_id']}: {row['delta']:.6f} nats/token" for row in largest)+". "
    spread=max(row["seed_delta_max"]-row["seed_delta_min"] for row in paired["per_book"])
    text+=f"Maximum per-book delta range across actual seeds: {spread:.12g}. "
    inputs=read_optional("results/final_input_manifest.json")
    choices={}; fallback=0
    for relative in inputs["improvement"]["candidate"]:
        result=read_optional(relative)
        for book in result["per_book"]:
            selection=book["selection"]
            label=str(selection["sink_tokens"]) if selected=="h1_adaptive_sink" else str(selection["anchors"])
            choices[label]=choices.get(label,0)+1; fallback+=int(selection["fallback"])
    text+=("Selected k counts" if selected=="h1_adaptive_sink" else "Frozen anchor-set counts")+" across the 30 actual book-seed evaluations: "+str(choices)+f"; fallbacks {fallback}. "
    if selected=="h1_adaptive_sink" and set(choices)=={"4"}:
        text+="H1 retained the baseline's four sinks in every evaluation; a quality gain cannot be attributed to adaptation in this case. "
    text+="These are descriptive cases on the frozen test set, not a further tuning rule."
    return text


def position_text(stage):
    path=ROOT/"figures"/(stage+"_nll_vs_position.csv")
    if not path.exists(): return "Position curves await complete paired outputs."
    with path.open() as stream: rows=list(csv.DictReader(stream))
    curves={label:{int(row["first_step"]):row for row in rows if row["method"]==label} for label in ("baseline","candidate")}
    common=sorted(set(curves["baseline"])&set(curves["candidate"]))
    values=[]
    for step in (common[0],common[-1]):
        base,cand=curves["baseline"][step],curves["candidate"][step]
        values.append(f"bin starting {step}: candidate-minus-baseline {float(cand['mean_book_bin_nll'])-float(base['mean_book_bin_nll']):.6f} nats/token, {base['n_books']} books")
    return "Descriptive 512-step equal-book position bins: "+"; ".join(values)+". Partial last bins are not zero-padded; exact book coverage is saved in the companion CSV. Bin differences are secondary and do not replace the primary whole-scored-region book estimand."


def sections():
    diagnostic=json.loads((ROOT/"results/diagnostics/summary.json").read_text())
    completion=json.loads((ROOT/"results/completion.json").read_text())
    rows=diagnostic["runs"]
    diag="; ".join(f'{r["run_id"]}: NLL {r["macro_book_nll"]:.6f}' for r in rows)
    reproduction=read_optional("results/reproduction/summary.json")
    improvement=read_optional("results/improvement/summary.json")
    reference=read_optional("environment/verification/reference-error-distributions.json")
    reference_status="Measured tiny-fixture GPU error distributions are recorded separately; they do not bound all full-model errors." if reference and reference.get("status")=="pass" else "The preset-tolerance GPU error-distribution capture is pending as a separate engineering release check."
    decisions=json.loads((ROOT/"docs/approval_decisions.json").read_text())
    selected=decisions["selected_improvement"].get("method")
    pilot_rows=[]
    if (ROOT/"results/pilots/comparison.csv").exists():
        with open(ROOT/"results/pilots/comparison.csv") as stream: pilot_rows=list(csv.DictReader(stream))
    pilots="; ".join(f"{row['method']}: development delta {float(row['paired_macro_delta_nll']):.6f}, 95% CI [{float(row['ci95_low']):.6f}, {float(row['ci95_high']):.6f}], runtime ratio {float(row['runtime_ratio']):.4f}, peak-memory ratio {float(row['peak_memory_ratio']):.4f}, acceptable cost {row['cost_acceptable']}" for row in pilot_rows) or "Registered development pilots are still running."
    controls=[]
    for name,state in completion.get("ablations",{}).items():
        value=read_optional("results/ablations/"+name+"/summary.json")
        if value: controls.append(name+": "+paired_text(value))
    control_text=" ".join(controls) or "Selected controls are pending complete raw outputs."
    identity=read_optional("environment/verification/forced-prefix-identity.json")
    if identity and identity.get("status")=="pass":
        control_text+=f" Full-checkpoint forced-prefix control: all {len(identity['comparisons'])} development book-seed NLL/input/mask arrays match the fixed-four baseline bitwise; exact receipt environment/verification/forced-prefix-identity.json."
    core_ready=all(completion.get(stage,{}).get("state")=="complete" for stage in ("reproduction","pilots","improvement")) and len(completion.get("ablations",{}))==2 and all(v.get("state")=="complete" for v in completion["ablations"].values())
    evidence_status="The registered scientific execution is complete and source-linked; real human course activities are separately recorded." if core_ready else "Scientific execution is authorized and in progress; pending stages have no assigned final verdict."
    return [
        ("Claim and current evidence", "We study whether retaining four initial attention-sink tokens improves equal-book post-overflow NLL over window attention with a 1024-position retained KV budget. "+evidence_status+(" Main result: "+paired_text(reproduction)+" Selected-method result: "+paired_text(improvement) if core_ready else "")+" The user explicitly waived procedural approvals and the 20 GPU-hour ceiling on 2026-10-09. This is an owner instruction, not a teacher approval or a peer signature. The main claim uses complete strict paired outputs, never a historical smoke."),
        ("Mechanism and position correction", "StreamingLLM retains initial sinks and recent K/V. With GPT-NeoX RoPE, raw keys must be cached before rotation and receive positions continuous within the current cache on each forward. Matching only the new token's position ID cannot establish correct query-key distances. The pinned official patch is enabled on every layer of both arms. This restores the published baseline mechanism and is not an original improvement."),
        ("Fixed data and scoring", "Pinned Pythia-2.8B and tokenizer; original ten PG19 test books, independent book resets. Input cap 16384; book 12204 has 7141 tokens. Total inputs 154597, predictions 154587, scored positions 144337 per arm/seed. Input x[i] predicts x[i+1]; eviction follows forward. First scored step 1025 predicts target x[1026]. A 4096-token smoke therefore scores 3070 positions. Text/model/token hashes and all per-book masks are checked."),
        ("Statistics and independent verification", "Primary estimand: seed-paired per-book NLL deltas averaged within book, then averaged equally over books. Percentile 95% CI uses 10000 book-cluster resamples, RNG seed 0. Actual model seeds 0, 1, 2 are distinct from bootstrap seeds and never create 30 independent books. Micro token NLL/PPL is secondary; exp(macro NLL) is macro-derived PPL. Fixtures test known probabilities, unequal lengths, boundaries and missing/duplicate/nonfinite contracts. An independent three-layer tiny GPT-NeoX reference checks trigonometric RoPE, raw-K, every layer K/V, attention and final logits through multiple evictions; real-model integration uses the full 32-layer pinned checkpoint. H800 fp16/bf16 tolerances are preset. "+reference_status+" The audit script reconstructs NLL, PPL, book deltas, bootstrap intervals and actual seed spread directly from saved NPZ arrays; this remains agent engineering verification."),
        ("Measured 2 by 2 diagnostic", diag+". Each run uses one historically exposed smoke book, 4096 inputs and 3070 scored positions. Faithful positions change both absolute losses and the method gap. Actual retained K+V bytes plateau 480MiB in legacy and 320MiB in faithful positions, matching between methods within each position policy. This local effect diagnoses the old implementation; it cannot establish a ten-book population conclusion, systems speedup or million-token stability. Position NLL arrays, KV traces, launch source hashes and costs are preserved."),
        ("Formal reproduction and paper comparison", paired_text(reproduction)+" "+position_text("reproduction")+" Paper arXiv v4 section 3.2 defines cache-relative positions; Figure 3 gives short-stream comparisons, Figure 5 super-long-stream results, and Figure 4 the cache schematic. The paper concatenates books; independent-book reset and post-overflow aggregation are explicit deviations. Formal methods are paired on the same existing H800 MIG 2g.20gb instance for each seed, with the same 30-SM hardware class across seeds. Full-GPU/MIG fp16 diagnostic outputs were not bitwise identical; the scored smoke-prefix mean shift was 0.000489 nats/token. Their results are not mixed in a formal pair. No paper speedup or full-H800 throughput equivalence is inferred."),
        ("Three improvement hypotheses and two pilots", "H1 selects the smallest k in {1,2,4,8} covering 90% of first-eight attention mass from queries 64 through 511, then freezes k within the book; negligible/invalid mass falls back to four. H2 protects position 0 and chooses the three highest-mass positions 1 through 63, breaks ties earlier, preserves chronological order and deduplicates recent K/V. Both use only causal prefix attention and a fixed 1024 retained budget. H3 is optional preallocated rolling-KV systems work, motivated only if a profiler shows material allocation cost; it is not implemented. Validation books 1022, 11155 and 13089 were selected without losses; cap 8192 and actual seeds 0/1/2 were frozen before pilot evaluation."),
        ("Pilots, selection, full evaluation and controls", pilots+". Selection follows the frozen lower-mean-NLL/cost rule with runtime and peak-memory ratios at most 1.10. If neither qualifies, simpler H1 is evaluated as the preregistered negative-result fallback. The owner delegated this rule; no parameter search or test selection is added. Recorded selection: "+str(selected)+". Full comparison against fixed-four StreamingLLM: "+paired_text(improvement)+" Selected controls: "+control_text+" Forced-prefix identity is independently tested; H1 has calibration-only/fixed-k controls, H2 forced-prefix/random-anchor controls. Controls use the development set and cannot be treated as full-test effects. Only the selected family's controls are required and executed. "+selected_cases(selected)+" Per-book frozen anchors/k and calibration intervals are in tables/improvement_book_selections.csv."),
        ("Compute, provenance and reproducibility", f"Current spent/reserved charge is {spent():.6f} device-instance hours, with a conservative 0.27h historical reserve. Full-GPU and MIG instance wall-hours are not normalized billing or currency cost. The user authorized unrestricted resources; 20h is no longer an execution limit. Three existing idle H800 MIG instances run seed workers concurrently, preserving another project on GPU0. Each attempt records selected UUID/parent, 30-SM class, driver, clock/power snapshot, source/config/data hashes, status, per-book recovery and registered-run wall time. The timer begins after the launch Git snapshot and run-directory creation and ends at finish(), includes model loading/inference/failure intervals, and excludes preflight/startup, the initial Git snapshot and post-finish checksum/exit tails. These are the same stored runtime_seconds used by the frozen cost rule; no cost field or threshold is changed. Runtime comparisons use the matched MIG class; clocks are not locked. Calibration intervals include necessary prefix forward passes and attention collection, measured as host wall time; they are not isolated extra kernel costs and are not added again to the registered-run charge. Incremental cost is assessed through registered-run ratios and controls. Pinned Python 3.10.12 / torch 2.14.0+cu130 / Transformers 4.33.0 and NVML bindings restore through setup. GPU1 instances are unaffected by GPU0's nightly cron."),
        ("Limitations, AI reflection and human delivery", "The ten fixed books are not a random sample of all PG19; one test book was previously observed. Three deterministic evaluations do not add independent books. Development selection uses only three books; attention proxies may fail, quality may not improve and calibration may exceed cost thresholds. Runtime order is not counterbalanced, clocks are unlocked and the three instances share a physical parent; ratios support the registered descriptive cost gate, not a causal speedup claim. The report preserves negative, inconclusive and failed outcomes. Agent verification is engineering verification, not an independent peer audit. The user's explicit authorization supersedes procedural approvals, while no teacher approval, member contribution, critical human review, deadline, final template or submission receipt is fabricated. Actual roster, T_final, peer review and individual defense remain human facts. Members can use the peer/defense packet. Exact Codex serving version was not exposed."),
        ("References and artifact entry points", "StreamingLLM: https://arxiv.org/html/2309.17453v4 ; official code pinned in third_party/streaming-llm. Model: pinned EleutherAI/pythia-2.8b revision and Apache-2.0 model card. PG19: DeepMind official processed GCS texts and pinned split lists. Rebuild: python scripts/build_results.py --manifest results/final_input_manifest.json ; python scripts/build_deliverables.py. Input/run hashes: results/build_provenance.json. Decisions: docs/review_packet.md. AI/roles: AI_USAGE.md and CONTRIBUTIONS.md."),
    ]


def metric_bullets(value,pending):
    if not value: return [pending,"No final numerical verdict is assigned before all registered outputs validate."]
    low,high=value["bootstrap_ci95"]
    return [f"Paired equal-book delta {value['paired_macro_delta_nll']:.6f} nats/token; 95% CI [{low:.6f}, {high:.6f}]. Verdict: {value['claim_status']}.",
            f"Macro NLL {value['baseline']['macro_book_nll']:.6f} → {value['candidate']['macro_book_nll']:.6f}; secondary micro PPL {value['baseline']['micro_token_ppl']:.3f} → {value['candidate']['micro_token_ppl']:.3f}.",
            f"{value['n_books']} books, actual seeds {value['seeds']}; seeds are averaged within each book before the book bootstrap."]


def deck_specs(content,ready):
    reproduction=read_optional("results/reproduction/summary.json")
    improvement=read_optional("results/improvement/summary.json")
    decision=read_optional("docs/approval_decisions.json")["selected_improvement"]
    selected=decision.get("method")
    pilot_lines=[]
    path=ROOT/"results/pilots/comparison.csv"
    if path.exists():
        with path.open() as stream:
            for row in csv.DictReader(stream):
                pilot_lines.append(f"{row['method']}: delta {float(row['paired_macro_delta_nll']):.6f}; runtime {float(row['runtime_ratio']):.3f}x; peak memory {float(row['peak_memory_ratio']):.3f}x; cost acceptable: {row['cost_acceptable']}.")
    controls=[]
    completion=read_optional("results/completion.json")
    for name in completion.get("ablations",{}):
        summary=read_optional("results/ablations/"+name+"/summary.json")
        if summary:
            low,high=summary["bootstrap_ci95"]
            controls.append(f"{name}: delta {summary['paired_macro_delta_nll']:.6f}; CI [{low:.6f}, {high:.6f}]; {summary['claim_status']}.")
    def slide(title,bullets,note,figure=None):
        return dict(title=title,bullets=bullets,notes=note,figure=figure)
    return [
        slide("StreamingLLM on H800",["Measured scientific delivery" if ready else "Registered scientific execution in progress",
              "Main comparison: fixed-four StreamingLLM versus window attention.","Selected improvement: "+str(selected)+"; matched H800 MIG 2g.20gb instances."],content[0][1]),
        slide("Why positions matter",["Retain initial sinks and recent K/V within a fixed budget.",
              "Cache raw K before RoPE; rotate every layer using positions continuous within the current cache.",
              "A current-token position ID alone cannot verify historical query-key distances.","The position repair restores the baseline; H1/H2 are separate improvement hypotheses."],content[1][1]),
        slide("Frozen inputs and scoring",["Pinned Pythia-2.8B; original ten PG19 test books; per-book cache reset; cap 16,384.",
              "Input x[i] predicts x[i+1]; evict after forward. Score steps ≥1,025, first target x[1,026].",
              "Retain ≤1,024 positions; forward may see 1,025. Total 144,337 scored positions per arm/seed.",
              "Equal-book paired NLL is primary; pooled token PPL is secondary."],content[2][1]),
        slide("Independent correctness evidence",["Known-probability scorer fixtures and unequal-book-length aggregation counterexamples.",
              "Three-layer tiny-model oracle: every-layer trigonometric RoPE, raw-K/V, attention and final logits; full-checkpoint integration separately.",
              "Strict book/token/mask/source/device contracts; missing, duplicate and nonfinite outputs fail.",
              "57 CPU checks and 2 actual H800 GPU precision checks pass; failed attempts remain visible."],content[3][1]),
        slide("Formal reproduction",metric_bullets(reproduction,"Six registered main runs are in progress."),content[5][1],"figures/reproduction_per_book_delta.png"),
        slide("What the old smoke concealed",["Legacy: window 5.450578, sinks 5.292458; local delta −0.158119.",
              "Corrected positions: window 3.051129, sinks 2.223681; local delta −0.827448.",
              "One exposed book / 3,070 scored positions: mechanism diagnosis, separate from the formal verdict."],content[4][1],"figures/diagnostic_nll_vs_position.png"),
        slide("Three falsifiable hypotheses",["H1: freeze the smallest k in {1,2,4,8} covering 90% of first-eight causal prefix attention mass.",
              "H2: keep position 0 plus three causal attention-ranked anchors among positions 1–63.",
              "H3: optional preallocated rolling KV, conditional on profiler evidence; not implemented.",
              "No test tuning. Two validation pilots use three frozen books, 8,192-token caps and seeds 0/1/2."],content[6][1]),
        slide("Pilots and frozen selection",pilot_lines+["Recorded selection: "+str(selected)+". Lower negative mean delta wins subject to runtime/memory ≤1.10x; H1 is the negative-result fallback."],content[7][1]),
        slide("Full selected-method evaluation",metric_bullets(improvement,"Full evaluation follows the frozen pilot rule.")[:2]+([f"Registered-run runtime {improvement['candidate']['runtime_seconds']/improvement['baseline']['runtime_seconds']:.3f}x; peak allocated memory {improvement['candidate']['peak_gpu_memory_mb']/improvement['baseline']['peak_gpu_memory_mb']:.3f}x."] if improvement else []),content[7][1],"figures/improvement_quality_cost.png"),
        slide("Controls and failure cases",controls+["Controls use the three development books. Forced-prefix/calibration-only identities are independently checked.",
              "Preserve negative/inconclusive outcomes, original raw traces and per-book selections; no further test search."],content[7][1]+"\n\nPer-book cases: results/improvement/per_book.csv and tables/improvement_book_selections.csv."),
        slide("Reproduce the evidence",["New Python 3.10 environment: all 67 locked dependency versions match; 57 CPU checks pass.",
              "Public model/tokenizer/text recovery and all 13 frozen-book hashes validate.",
              "Immutable run → config/source/data → NPZ → metric → table/figure hash chain.",
              "Actual terminal demo and source-linked fallback; owner authorization and AI actions are recorded."],content[8][1]+"\n\n"+content[10][1]),
        slide("Limits and defensible conclusions",["Ten fixed books and one historically exposed smoke book; CI does not establish all-PG19 generality.",
              "Independent resets, 16K streams and MIG hardware limit direct comparisons with the paper.",
              "No million-token, training or speedup claim. Runtime is measured with unlocked clocks.",
              "Scientific/engineering evidence is complete only when validated. Peer review, real contributions and individual defense remain human facts."],content[9][1]),
    ]


def markdown_tables():
    groups=[("Reproduction per book","tables/reproduction_per_book.csv",["book_id","baseline_nll","candidate_nll","delta","scored_tokens"]),
            ("Development pilots","tables/pilot_comparison.csv",["method","paired_macro_delta_nll","ci95_low","ci95_high","runtime_ratio","peak_memory_ratio"]),
            ("Selected-method full test","tables/improvement_per_book.csv",["book_id","baseline_nll","candidate_nll","delta","scored_tokens"]),
            ("Scientific summary and controls","tables/scientific_summary.csv",["comparison","n_books","paired_macro_delta_nll","ci95_low","ci95_high","claim_status","runtime_ratio","peak_memory_ratio"])]
    blocks=[]
    for title,relative,columns in groups:
        path=ROOT/relative
        if not path.exists(): continue
        with path.open() as stream: rows=list(csv.DictReader(stream))
        def formatted(value):
            try: return f"{float(value):.6f}" if any(character in value for character in (".","e","E")) else value
            except ValueError: return value
        table="| "+" | ".join(columns)+" |\n|"+"|".join("---" for _ in columns)+"|\n"
        table+="\n".join("| "+" | ".join(formatted(row[column]) for column in columns)+" |" for row in rows)
        blocks.append("## "+title+"\n\nSource: `"+relative+"`. Negative delta favors the candidate; controls use development books.\n\n"+table)
    return "\n\n".join(blocks)


def main():
    p=argparse.ArgumentParser(description=__doc__); p.parse_args()
    (ROOT/"report").mkdir(exist_ok=True); (ROOT/"presentation").mkdir(exist_ok=True)
    content=sections()
    completion=json.loads((ROOT/"results/completion.json").read_text())
    ready=all(completion.get(stage,{}).get("state")=="complete" for stage in ("reproduction","pilots","improvement")) and len(completion.get("ablations",{}))==2 and all(v.get("state")=="complete" for v in completion["ablations"].values())
    report_title="StreamingLLM: measured H800 scientific delivery" if ready else "StreamingLLM: execution progress report"
    subtitle="Owner-authorized complete scientific evidence; human course activities separately recorded." if ready else "Registered execution in progress; incomplete stages have no assigned verdict."
    stem="report" if ready else "report_draft"
    deck_name="defense.pptx" if ready else "defense_draft.pptx"
    md="# "+report_title+"\n\n"+subtitle+"\n\n"+"\n\n".join("## "+section+"\n\n"+body for section,body in content)+"\n\n"+markdown_tables()+"\n"
    (ROOT/("report/"+stem+".md")).write_text(md.rstrip()+"\n")
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle
    styles=getSampleStyleSheet(); story=[Paragraph(report_title,styles["Title"]),Paragraph(subtitle,styles["Normal"]),Spacer(1,12)]
    def csv_table(relative,columns):
        path=ROOT/relative
        if not path.exists(): return []
        with path.open() as stream: rows=list(csv.DictReader(stream))
        data=[[label for _,label,_ in columns]]
        for row in rows:
            data.append([formatter(row[key]) for key,_,formatter in columns])
        table=Table(data,repeatRows=1,hAlign="LEFT")
        table.setStyle(TableStyle([("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("FONTSIZE",(0,0),(-1,-1),8),
                                  ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#e8eef5")),("GRID",(0,0),(-1,-1),.3,colors.lightgrey),
                                  ("BOTTOMPADDING",(0,0),(-1,-1),5),("TOPPADDING",(0,0),(-1,-1),5)]))
        return [table,Spacer(1,10)]
    number=lambda value:f"{float(value):.6f}"
    for i,(title,body) in enumerate(content):
        story.extend([Paragraph(title,styles["Heading2"]),Paragraph(html.escape(body),styles["Normal"]),Spacer(1,8)])
        if i==4: story.append(Image(str(ROOT/"figures/diagnostic_nll_vs_position.png"),width=480,height=240))
        if i==5 and (ROOT/"figures/reproduction_per_book_delta.png").exists(): story.append(Image(str(ROOT/"figures/reproduction_per_book_delta.png"),width=480,height=240))
        if i==5:
            story.extend(csv_table("tables/reproduction_per_book.csv",[("book_id","Book",str),("baseline_nll","Window NLL",number),("candidate_nll","Sink NLL",number),("delta","Paired delta",number),("scored_tokens","Scored",str)]))
            if (ROOT/"figures/reproduction_cache.png").exists(): story.append(Image(str(ROOT/"figures/reproduction_cache.png"),width=480,height=187))
        if i==7 and (ROOT/"figures/improvement_per_book_delta.png").exists(): story.append(Image(str(ROOT/"figures/improvement_per_book_delta.png"),width=480,height=240))
        if i==7:
            story.extend(csv_table("tables/pilot_comparison.csv",[("method","Pilot",str),("paired_macro_delta_nll","Delta NLL",number),("runtime_ratio","Runtime ratio",number),("peak_memory_ratio","Memory ratio",number)]))
            story.extend(csv_table("tables/improvement_per_book.csv",[("book_id","Book",str),("baseline_nll","Fixed-four NLL",number),("candidate_nll","Selected NLL",number),("delta","Paired delta",number),("scored_tokens","Scored",str)]))
            if (ROOT/"figures/improvement_quality_cost.png").exists(): story.append(Image(str(ROOT/"figures/improvement_quality_cost.png"),width=480,height=192))
    def footer(canvas,doc):
        canvas.saveState(); canvas.setFont("Helvetica",8); canvas.drawString(40,25,"MEASURED SCIENTIFIC DELIVERY | HUMAN COURSE ACTIVITIES SEPARATELY RECORDED" if ready else "EXECUTION PROGRESS | NO VERDICT FOR INCOMPLETE STAGES"); canvas.drawRightString(555,25,str(doc.page)); canvas.restoreState()
    SimpleDocTemplate(str(ROOT/("report/"+stem+".pdf")),title=report_title,author="CISC8006 project; agent-assisted evidence",pagesize=A4,leftMargin=40,rightMargin=40,topMargin=40,bottomMargin=42).build(story,onFirstPage=footer,onLaterPages=footer)
    from pptx import Presentation
    from pptx.util import Inches, Pt
    presentation=Presentation(); presentation.slide_width=Inches(13.333); presentation.slide_height=Inches(7.5)
    slide_specs=deck_specs(content,ready)
    for spec in slide_specs:
        slide=presentation.slides.add_slide(presentation.slide_layouts[6])
        title_box=slide.shapes.add_textbox(Inches(.65),Inches(.45),Inches(12),Inches(.8)); title_box.text=spec["title"]
        title_box.text_frame.paragraphs[0].font.size=Pt(30)
        figure=ROOT/spec["figure"] if spec.get("figure") else None
        has_figure=figure is not None and figure.is_file()
        box=slide.shapes.add_textbox(Inches(.7),Inches(1.45),Inches(11.9),Inches(2.5 if has_figure else 5.3)); frame=box.text_frame; frame.word_wrap=True
        for j,line in enumerate(spec["bullets"]):
            paragraph=frame.paragraphs[0] if j==0 else frame.add_paragraph(); paragraph.text=line; paragraph.font.size=Pt(18 if has_figure else 23); paragraph.space_after=Pt(10)
        if has_figure:
            from PIL import Image as PILImage
            with PILImage.open(figure) as picture: aspect=picture.width/picture.height
            height=2.65; width=height*aspect
            slide.shapes.add_picture(str(figure),Inches((13.333-width)/2),Inches(4.15),width=Inches(width),height=Inches(height))
        slide.notes_slide.notes_text_frame.text=spec["notes"]
        footer_box=slide.shapes.add_textbox(Inches(.7),Inches(7.02),Inches(12),Inches(.28)); footer_box.text="Exact inputs in results/build_provenance.json | owner-authorized execution | human activities separately recorded"; footer_box.text_frame.paragraphs[0].font.size=Pt(10)
    presentation.save(ROOT/("presentation/"+deck_name))
    outputs=[ROOT/("report/"+stem+".md"),ROOT/("report/"+stem+".pdf"),ROOT/("presentation/"+deck_name)]
    write_json(ROOT/"submission/deliverables_manifest.json",dict(built_at=now(),state="scientific_complete_owner_authorized" if ready else "execution_progress",command=[sys.executable]+sys.argv,builder_sha256=file_hash(Path(__file__)),artifacts=[dict(path=str(path.relative_to(ROOT)),sha256=file_hash(path)) for path in outputs],input_manifest_sha256=file_hash(ROOT/"results/final_input_manifest.json")))
    print("Built "+", ".join(str(path.relative_to(ROOT)) for path in outputs))


if __name__=="__main__": main()
